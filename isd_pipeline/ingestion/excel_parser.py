"""Parse one Power BI Excel report into standardized snapshot rows (schema 3a).

What it does, in order:
1. Find the real header row (scan the top N rows; pick the row matching the most aliases).
2. Map headers to canonical names via config/report_schema.yaml (+ per-division overrides).
   If the header repeats (a second side-by-side copy of the table), use only the first block.
3. Walk the data rows. Inspector names are filled down ONLY when the sheet structure
   justifies it: (a) merged inspector cells, or (b) an inspector group-header row (inspector
   only, no other values) with blank inspector cells beneath it, until a blank or total row.
   No other column is ever filled down.
4. Read is_red from the Call Date cell's font colour (RGB, indexed and theme colours; theme
   accent colours are ambiguous -> not red, logged).
5. Validate. Missing required columns -> the whole file is invalid (no rows returned).
   Row-level problems are logged as invalid rows; the row is kept (flagged) - never dropped.
"""
from __future__ import annotations

import io
import logging
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

from isd_pipeline.schema import SNAPSHOT_COLUMNS, coerce
from isd_pipeline.settings import division_schema, load_settings

log = logging.getLogger(__name__)

KEY_ROW_FIELDS = ("ap_number", "call_datetime", "address", "insp_type")


@dataclass
class ValidationResult:
    division: str
    report_date: date
    is_valid: bool = True
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    sheet_name: str = ""
    header_row: int | None = None
    header_map: dict[str, str] = field(default_factory=dict)   # canonical -> "C (Call Date)"
    primary_block: str = ""
    secondary_block_detected: bool = False
    rows_extracted: int = 0
    invalid_rows: list[dict] = field(default_factory=list)
    filled_inspector_cells: int = 0
    fill_log: list[str] = field(default_factory=list)
    unfilled_blank_inspector_rows: list[int] = field(default_factory=list)
    skipped_rows: list[str] = field(default_factory=list)
    red_cells_checked: str = ""
    red_count: int = 0
    ambiguous_colors: list[str] = field(default_factory=list)

    @property
    def n_invalid_rows(self) -> int:
        return len(self.invalid_rows)

    def summary(self) -> dict[str, Any]:
        d = asdict(self)
        d["report_date"] = self.report_date.isoformat()
        d["n_invalid_rows"] = self.n_invalid_rows
        return d


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------
def normalize_header(text: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _text(v: Any) -> str:
    """Cell -> text without losing leading zeros (strings are kept verbatim)."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.isoformat(sep=" ")
    return str(v)


def classify_font_color(color: Any) -> tuple[bool, str | None]:
    """Return (is_red, ambiguity_reason). Ambiguous colours are reported as not red."""
    if color is None:
        return False, None
    ctype = getattr(color, "type", None)
    rgb: str | None = None
    if ctype == "rgb":
        rgb = color.rgb if isinstance(color.rgb, str) else None
        if rgb is None:
            return False, "unreadable rgb value"
    elif ctype == "indexed":
        idx = color.indexed
        if isinstance(idx, int) and 0 <= idx < 64 and idx < len(COLOR_INDEX):
            rgb = COLOR_INDEX[idx]
        else:
            return False, f"system indexed colour {idx}"
    elif ctype == "theme":
        # Theme 0-3 are the text/background slots (black/white by default); 4-9 are accents
        # whose actual colour depends on the workbook theme.
        if color.theme in (0, 1, 2, 3):
            return False, None
        return False, f"theme colour {color.theme} (tint {color.tint})"
    elif ctype == "auto":
        return False, None
    else:
        return False, f"unknown colour type {ctype}"
    rgb = rgb[-6:]
    try:
        r, g, b = int(rgb[0:2], 16), int(rgb[2:4], 16), int(rgb[4:6], 16)
    except ValueError:
        return False, f"unparseable rgb {rgb}"
    return (r >= 0xC0 and g <= 0x60 and b <= 0x60), None


def parse_call_datetime(value: Any, formats: list[str]) -> tuple[datetime | None, str | None]:
    """Return (timestamp, problem). Time of day is always kept."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "missing_call_datetime"
    if isinstance(value, datetime):
        return value.replace(tzinfo=None), None
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day), None
    if isinstance(value, (int, float)):
        try:
            return from_excel(value), None
        except Exception:  # noqa: BLE001
            return None, f"unparseable_call_datetime: {value!r}"
    text = str(value).strip()
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt), None
        except ValueError:
            continue
    ts = pd.to_datetime(text, errors="coerce")
    if pd.isna(ts):
        return None, f"unparseable_call_datetime: {text!r}"
    return ts.to_pydatetime().replace(tzinfo=None), None


# ---------------------------------------------------------------------------
# Header detection
# ---------------------------------------------------------------------------
def _alias_lookup(schema: dict[str, Any]) -> dict[str, str]:
    lookup: dict[str, str] = {}
    for canon, aliases in (schema.get("aliases") or {}).items():
        for alias in [canon, *aliases]:
            lookup.setdefault(normalize_header(alias), canon)
    return lookup


def detect_header_row(ws, lookup: dict[str, str], scan_rows: int) -> tuple[int | None, int]:
    best_row, best_score = None, 0
    for r in range(1, min(scan_rows, ws.max_row) + 1):
        matched = set()
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str):
                canon = lookup.get(normalize_header(v))
                if canon:
                    matched.add(canon)
        if len(matched) > best_score:
            best_row, best_score = r, len(matched)
    return best_row, best_score


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------
def process_excel_report(
    file_path: str | Path | bytes,
    division: str,
    report_date: date,
    *,
    source: dict[str, str] | None = None,
    batch_id: str = "",
    schema: dict[str, Any] | None = None,
    scan_rows: int | None = None,
    min_header_matches: int | None = None,
    ingested_at: datetime | None = None,
) -> tuple[pd.DataFrame, ValidationResult]:
    """Parse one report. Returns (snapshot rows, ValidationResult).

    file_path may be a path or the raw bytes of the attachment. `source` carries
    email_id / attachment_id / filename for provenance columns.
    """
    if schema is None or scan_rows is None or min_header_matches is None:
        settings = load_settings()
        schema = schema or settings.report_schema
        scan_rows = scan_rows or settings.header_scan_rows
        min_header_matches = min_header_matches or settings.min_header_matches
    dschema = division_schema(schema, division)
    source = source or {}
    vr = ValidationResult(division=division, report_date=report_date)
    empty = coerce(pd.DataFrame(), SNAPSHOT_COLUMNS)

    try:
        data = file_path if isinstance(file_path, (bytes, bytearray)) else Path(file_path).read_bytes()
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001
        vr.is_valid = False
        vr.errors.append(f"Could not open the Excel file: {exc}")
        return empty, vr

    lookup = _alias_lookup(dschema)
    # Pick the sheet whose best header row matches the most columns.
    best = None
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        row, score = detect_header_row(ws, lookup, scan_rows)
        if row and (best is None or score > best[2]):
            best = (ws, row, score)
    if best is None or best[2] < min_header_matches:
        vr.is_valid = False
        found = best[2] if best else 0
        vr.errors.append(f"No header row found in the top {scan_rows} rows "
                         f"(best match {found} known columns, need {min_header_matches})")
        return empty, vr
    ws, header_row, _score = best
    vr.sheet_name, vr.header_row = ws.title, header_row

    # --- map columns, detect a second side-by-side block ---
    colmap: dict[str, int] = {}
    header_text: dict[int, str] = {}
    block_end = ws.max_column
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if not isinstance(v, str) or not v.strip():
            continue
        canon = lookup.get(normalize_header(v))
        if canon is None:
            if not colmap or c <= block_end:
                vr.warnings.append(f"Unrecognised column '{v}' at {get_column_letter(c)}{header_row} ignored")
            continue
        if canon in colmap:
            block_end = c - 1
            vr.secondary_block_detected = True
            vr.warnings.append(
                f"Second side-by-side copy of the table starts at column {get_column_letter(c)}; "
                f"using only the primary block")
            break
        colmap[canon] = c
        header_text[c] = v.strip()
    first_col = min(colmap.values())
    last_col = max(colmap.values())
    vr.primary_block = f"{get_column_letter(first_col)}:{get_column_letter(last_col)}"
    vr.header_map = {k: f"{get_column_letter(c)} ({header_text[c]})" for k, c in colmap.items()}

    missing = [c for c in dschema.get("required_columns", []) if c not in colmap]
    if missing:
        vr.is_valid = False
        vr.errors.append(f"Missing required column(s): {', '.join(missing)} "
                         f"(header row {header_row}: {', '.join(header_text.values())})")
        return empty, vr

    # --- merged inspector cells ---
    insp_col = colmap["inspector"]
    merged_fill: dict[int, tuple[str, str]] = {}
    merged_top: set[int] = set()
    for rng in ws.merged_cells.ranges:
        if rng.min_col <= insp_col <= rng.max_col and rng.min_row > header_row:
            top = ws.cell(row=rng.min_row, column=rng.min_col).value
            if _blank(top):
                continue
            merged_top.add(rng.min_row)
            for r in range(rng.min_row + 1, rng.max_row + 1):
                merged_fill[r] = (str(top).strip(), str(rng.coord))

    red_field = dschema.get("red_font_column", "call_datetime")
    red_col = colmap.get(red_field)
    formats = list(dschema.get("call_datetime_formats") or [])
    total_prefixes = tuple(p.lower() for p in dschema.get("total_row_prefixes", ["total"]))
    now = ingested_at or datetime.now(timezone.utc)

    records: list[dict[str, Any]] = []
    group_name: str | None = None
    group_source: str | None = None   # 'group_header' | 'merged' | 'data_row'
    data_rows: list[int] = []

    for r in range(header_row + 1, ws.max_row + 1):
        vals = {canon: ws.cell(row=r, column=c).value for canon, c in colmap.items()}
        nonblank = [k for k, v in vals.items() if not _blank(v)]
        if not nonblank and r not in merged_fill:
            group_name, group_source = None, None          # blank row ends a group
            continue
        texts = [str(ws.cell(row=r, column=c).value).strip() for c in range(first_col, last_col + 1)
                 if isinstance(ws.cell(row=r, column=c).value, str)]
        if texts and texts[0].lower().startswith(total_prefixes) and _blank(vals.get("ap_number")) \
                and _blank(vals.get("call_datetime")):
            vr.skipped_rows.append(f"row {r}: total row '{texts[0]}'")
            group_name, group_source = None, None
            continue
        business = [k for k in nonblank if k != "inspector"]
        if not business and r not in merged_fill:
            group_name, group_source = str(vals["inspector"]).strip(), "group_header"
            vr.skipped_rows.append(f"row {r}: inspector group header '{group_name}'")
            continue

        # --- inspector (justified fill-down only) ---
        cell_ref = f"{get_column_letter(insp_col)}{r}"
        inspector = _text(vals["inspector"]).strip()
        if r in merged_fill:
            inspector = merged_fill[r][0]
            vr.filled_inspector_cells += 1
            vr.fill_log.append(f"{cell_ref} <- '{inspector}' (merged range {merged_fill[r][1]})")
            group_name, group_source = inspector, "merged"
        elif inspector:
            group_name = inspector
            group_source = "merged" if r in merged_top else "data_row"
        elif group_source == "group_header" and group_name:
            inspector = group_name
            vr.filled_inspector_cells += 1
            vr.fill_log.append(f"{cell_ref} <- '{inspector}' (group header above)")
        else:
            vr.unfilled_blank_inspector_rows.append(r)
            vr.warnings.append(f"{cell_ref}: inspector blank and no merged cell/group header above - left blank")

        # --- values ---
        issues: list[str] = []
        call_dt, problem = parse_call_datetime(vals.get("call_datetime"), formats)
        if problem:
            issues.append(problem)
        ap_raw = vals.get("ap_number")
        if isinstance(ap_raw, (int, float)) and not isinstance(ap_raw, bool):
            vr.warnings.append(f"{get_column_letter(colmap['ap_number'])}{r}: Ap # stored as a number in Excel "
                               f"({ap_raw}); leading zeros may already be lost in the source")
        ap = _text(ap_raw).strip()
        for f in ("ap_number", "address", "insp_type"):
            if _blank(vals.get(f)):
                issues.append(f"missing_{f}")

        is_red = False
        if red_col is not None:
            font = ws.cell(row=r, column=red_col).font
            is_red, amb = classify_font_color(font.color if font else None)
            if amb:
                vr.ambiguous_colors.append(f"{get_column_letter(red_col)}{r}: {amb} -> treated as not red")
        vr.red_count += int(is_red)

        rec = {
            "inspector": inspector,
            "ward": _text(vals.get("ward")).strip(),
            "call_datetime": call_dt,
            "ap_number": ap,
            "address": _text(vals.get("address")),
            "insp_type": _text(vals.get("insp_type")).strip(),
            "work_group": _text(vals.get("work_group")).strip() if "work_group" in colmap else "",
            "contact_name": _text(vals.get("contact_name")),
            "contact_info": _text(vals.get("contact_info")),
            "request_notes": _text(vals.get("request_notes")),
            "inspector_notes": _text(vals.get("inspector_notes")),
            "fulfilled_late": _text(vals.get("fulfilled_late")).strip(),
            "is_red": is_red,
            "division": division,
            "report_date": pd.Timestamp(report_date),
            "source_email_id": source.get("email_id", ""),
            "source_attachment_id": source.get("attachment_id", ""),
            "source_filename": source.get("filename", ""),
            "ingested_at": now,
            "source_row_number": r,
            "ingestion_batch_id": batch_id,
            "row_issues": ";".join(issues),
        }
        if issues:
            vr.invalid_rows.append({
                "source_row_number": r,
                "reason": ";".join(issues),
                "values": {k: _text(v) for k, v in vals.items()},
            })
        records.append(rec)
        data_rows.append(r)

    if red_col is not None and data_rows:
        L = get_column_letter(red_col)
        vr.red_cells_checked = (f"{L}{data_rows[0]}:{L}{data_rows[-1]} "
                                f"({header_text.get(red_col, red_field)}, {len(data_rows)} cells)")
    if vr.ambiguous_colors:
        vr.warnings.append(f"{len(vr.ambiguous_colors)} Call Date cell(s) had an ambiguous font colour "
                           f"(treated as not red)")
    if not records:
        vr.warnings.append("The report contains no data rows")

    df = coerce(pd.DataFrame.from_records(records), SNAPSHOT_COLUMNS)
    vr.rows_extracted = len(df)
    log.info("%s %s: header row %d, %d rows, %d inspector cells filled, %d invalid rows",
             division, report_date, header_row, len(df), vr.filled_inspector_cells, vr.n_invalid_rows)
    return df, vr
