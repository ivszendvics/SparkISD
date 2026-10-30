"""Excel parser tests on small workbooks generated in the test."""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.styles.colors import Color

from isd_pipeline.ingestion.excel_parser import classify_font_color, process_excel_report

HDR = ["Inspector", "Ward", "Call Date", "Ap #", "Address", "Insp Type", "Contact Name", "Contact Info",
       "Request Notes", "Inspector Notes"]
D = date(2026, 10, 6)


def _wb(path: Path, rows: list[list], title_rows: int = 2, merges: list[str] = (), fonts: dict = None) -> Path:
    wb = Workbook()
    ws = wb.active
    for i in range(title_rows):
        ws.cell(row=i + 1, column=1, value="Permit Inspection Request - Building Summary" if i == 0 else None)
    h = title_rows + 1
    for c, v in enumerate(HDR, 1):
        ws.cell(row=h, column=c, value=v)
    for r, vals in enumerate(rows, start=h + 1):
        for c, v in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=v)
    for m in merges:
        ws.merge_cells(m)
    for ref, font in (fonts or {}).items():
        ws[ref].font = font
    wb.save(path)
    return path


def row(insp, ap, call=datetime(2026, 10, 5, 9, 14, 33), ward="07", addr="1 Main St", typ="Final"):
    return [insp, ward, call, ap, addr, typ, "Name", "(617) 555-0100", None, None]


def test_header_detection_with_title_rows(tmp_path):
    p = _wb(tmp_path / "a.xlsx", [row("Ann", "ALT1")], title_rows=5)
    df, vr = process_excel_report(p, "Building", D)
    assert vr.is_valid and vr.header_row == 6 and len(df) == 1


def test_ap_number_keeps_leading_zeros_and_time_kept(tmp_path):
    p = _wb(tmp_path / "a.xlsx", [row("Ann", "0012345")])
    df, _ = process_excel_report(p, "Building", D)
    assert df.loc[0, "ap_number"] == "0012345"
    assert df.loc[0, "ward"] == "07"
    assert df.loc[0, "call_datetime"] == datetime(2026, 10, 5, 9, 14, 33)


def test_merged_inspector_cells_filled(tmp_path):
    rows = [row("Ann", "A1"), row(None, "A2"), row(None, "A3"), row("Bob", "B1")]
    p = _wb(tmp_path / "a.xlsx", rows, merges=["A4:A6"])
    df, vr = process_excel_report(p, "Building", D)
    assert list(df["inspector"]) == ["Ann", "Ann", "Ann", "Bob"]
    assert vr.filled_inspector_cells == 2


def test_group_header_rows_filled_until_blank_row(tmp_path):
    g = lambda name: [name] + [None] * 9  # noqa: E731 - inspector-only group header row
    rows = [g("Ann"), row(None, "A1"), row(None, "A2"), [None] * 10, g("Bob"), row(None, "B1")]
    df, vr = process_excel_report(_wb(tmp_path / "a.xlsx", rows), "Building", D)
    assert list(df["inspector"]) == ["Ann", "Ann", "Bob"]
    assert vr.filled_inspector_cells == 3
    assert len(df) == 3                                   # header rows are not data


def test_no_fill_down_without_structure(tmp_path):
    """Flat layout: a blank inspector under a normal data row stays blank; other fields never fill."""
    rows = [row("Ann", "A1"), row(None, "A2", ward=None, addr=None)]
    df, vr = process_excel_report(_wb(tmp_path / "a.xlsx", rows), "Building", D)
    assert df.loc[1, "inspector"] == "" and df.loc[1, "ward"] == "" and df.loc[1, "address"] == ""
    assert vr.filled_inspector_cells == 0 and vr.unfilled_blank_inspector_rows == [5]
    assert "missing_address" in df.loc[1, "row_issues"]


def test_red_font_read_from_call_date_cell(tmp_path):
    rows = [row("Ann", "A1"), row("Ann", "A2"), row("Ann", "A3"), row("Ann", "A4")]
    fonts = {"C4": Font(color="FFFF0000"), "C5": Font(color=Color(indexed=10)),
             "C6": Font(color=Color(theme=5)), "D7": Font(color="FFFF0000")}   # D7 is Ap #, not Call Date
    df, vr = process_excel_report(_wb(tmp_path / "a.xlsx", rows, fonts=fonts), "Building", D)
    assert list(df["is_red"]) == [True, True, False, False]
    assert len(vr.ambiguous_colors) == 1 and "C6" in vr.ambiguous_colors[0]
    assert vr.red_cells_checked.startswith("C4:C7")


def test_font_color_classifier():
    assert classify_font_color(Color(rgb="FFC00000"))[0] is True       # dark red
    assert classify_font_color(Color(rgb="FF000000")) == (False, None)
    assert classify_font_color(Color(theme=1)) == (False, None)       # default text colour
    assert classify_font_color(None) == (False, None)


def test_missing_required_columns_invalid(tmp_path):
    wb = Workbook()
    ws = wb.active
    for c, v in enumerate(["Inspector", "Ward", "Address", "Insp Type", "Contact Name"], 1):
        ws.cell(row=1, column=c, value=v)
    ws.cell(row=2, column=1, value="Ann")
    wb.save(tmp_path / "bad.xlsx")
    df, vr = process_excel_report(tmp_path / "bad.xlsx", "Building", D)
    assert not vr.is_valid and df.empty
    assert "call_datetime" in vr.errors[0] and "ap_number" in vr.errors[0]


def test_row_level_problems_logged_and_kept(tmp_path):
    rows = [row("Ann", "A1"), row("Ann", None), row("Ann", "A3", call="not a date")]
    df, vr = process_excel_report(_wb(tmp_path / "a.xlsx", rows), "Building", D)
    assert len(df) == 3 and vr.n_invalid_rows == 2
    reasons = [r["reason"] for r in vr.invalid_rows]
    assert reasons[0] == "missing_ap_number" and reasons[1].startswith("unparseable_call_datetime")


def test_text_dates_and_side_by_side_block(tmp_path):
    wb = Workbook()
    ws = wb.active
    for off in (0, 11):
        for c, v in enumerate(HDR, 1):
            ws.cell(row=1, column=c + off, value=v)
        vals = row("Ann", "A1" if off == 0 else "COPY", call="10/05/2026 02:30 PM")
        for c, v in enumerate(vals, 1):
            ws.cell(row=2, column=c + off, value=v)
    wb.save(tmp_path / "s.xlsx")
    df, vr = process_excel_report(tmp_path / "s.xlsx", "Building", D)
    assert vr.secondary_block_detected and list(df["ap_number"]) == ["A1"]
    assert df.loc[0, "call_datetime"] == datetime(2026, 10, 5, 14, 30)


def test_corrupt_file_invalid(tmp_path):
    p = tmp_path / "x.xlsx"
    p.write_bytes(b"not an excel file")
    df, vr = process_excel_report(p, "Building", D)
    assert not vr.is_valid and df.empty
