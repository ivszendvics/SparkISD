"""Deterministic synthetic inbox for the ISD demo and the end-to-end tests.

Writes one folder per report day:

    <out>/<YYYY-MM-DD>/<division_slug>.json   email metadata (sender, subject, 07:03 received)
    <out>/<YYYY-MM-DD>/<division_slug>.xlsx   the Power BI-style attachment
    <out>/storyline.json                      the scripted scenarios (used by the demo spotlight)

All names, addresses, phone numbers (555-01xx) and Ap # values are fictitious.

Usage:
    python demo/generate_synthetic.py                    # -> demo_data/scripted_inbox
    python demo/generate_synthetic.py --deliver-all      # also copy every day into the mock inbox
"""
from __future__ import annotations

import argparse
import hashlib
import json
import random
import shutil
import sys
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.colors import Color

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

DEFAULT_OUT = ROOT / "demo_data" / "scripted_inbox"
SEED = 539

# ---------------------------------------------------------------------------
# Storyline calendar
# ---------------------------------------------------------------------------
REPORT_DAYS: list[date] = [
    date(2026, 10, 5), date(2026, 10, 6), date(2026, 10, 7), date(2026, 10, 8), date(2026, 10, 9),
    date(2026, 10, 12),  # Indigenous Peoples' Day: Power BI still sends the automated email
    date(2026, 10, 13), date(2026, 10, 14), date(2026, 10, 15), date(2026, 10, 16),
    date(2026, 10, 19), date(2026, 10, 20),
]
DEMO_DAY = date(2026, 10, 19)            # "Preload history" processes everything before this
ELECTRIC_MISSING_DAY = date(2026, 10, 20)
MALFORMED_PLUMBING_DAY = date(2026, 10, 15)
DUPLICATE_BUILDING_EMAIL_DAY = date(2026, 10, 8)
SIDE_BY_SIDE_ELECTRIC_DAY = date(2026, 10, 7)
AMBIGUOUS_COLOR_DAY = date(2026, 10, 7)

DIVISIONS = ("Building", "Electric", "Plumbing/Gas")
SLUG = {"Building": "building", "Electric": "electric", "Plumbing/Gas": "plumbing_gas"}
SUBJECT = {d: f"Permit Inspection Request - {d} Summary" for d in DIVISIONS}
SENDER_NAME = "Microsoft Power BI"
SENDER_ADDRESS = "no-reply-powerbi@mock.invalid"   # clearly fake; the real one goes in email_rules.yaml
RECIPIENT = "isd-team-member@bu.edu"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

INSPECTORS = {
    "Building": ["Maria Delgado", "Kevin O'Brien", "Thomas Nguyen", "Aisha Williams", "Robert Sullivan", "Hannah Lee"],
    "Electric": ["Daniel Murphy", "Linh Tran", "Patrick Walsh", "Sofia Ramirez", "James Fitzgerald"],
    "Plumbing/Gas": ["Michael Brennan", "Grace Okafor", "Luis Fernandes", "Erin McCarthy", "Samuel Cohen"],
}
# Hannah Lee gets few requests so the dashboard shows a low_n inspector.
INSPECTOR_WEIGHTS = {"Hannah Lee": 0.15}

INSP_TYPES = {
    "Building": ["Rough Frame", "Final Building", "Foundation", "Insulation", "Certificate of Occupancy"],
    "Electric": ["Rough Wiring", "Final Electrical", "Service Upgrade", "Temporary Service"],
    "Plumbing/Gas": ["Rough Plumbing", "Final Plumbing", "Gas Pressure Test", "Gas Meter Release"],
}
STREETS = [
    ("Dorchester", "Avenue"), ("Blue Hill", "Avenue"), ("Centre", "Street"), ("Washington", "Street"),
    ("Tremont", "Street"), ("Commonwealth", "Avenue"), ("Hyde Park", "Avenue"), ("Columbia", "Road"),
    ("River", "Street"), ("Bowdoin", "Street"), ("Adams", "Street"), ("Hancock", "Street"),
    ("Geneva", "Avenue"), ("Warren", "Street"), ("Huntington", "Avenue"), ("Boylston", "Street"),
    ("Beacon", "Street"), ("Cambridge", "Street"), ("Saratoga", "Street"), ("Bennington", "Street"),
    ("Meridian", "Street"), ("Brighton", "Avenue"), ("Market", "Street"), ("Belgrade", "Avenue"),
    ("Cummins", "Highway"), ("Talbot", "Avenue"), ("Morton", "Street"), ("Walk Hill", "Street"),
    ("Savin Hill", "Avenue"), ("Columbus", "Avenue"), ("Shawmut", "Avenue"), ("Harrison", "Avenue"),
]
ABBR = {"Avenue": "Ave", "Street": "St", "Road": "Rd", "Highway": "Hwy"}
FIRST = ["James", "Mary", "John", "Patricia", "Carlos", "Mei", "Ahmed", "Fatima", "Sean", "Brigid",
         "Tuan", "Rosa", "David", "Keisha", "Paulo", "Anna", "Marcus", "Yolanda", "Dmitri", "Chloe"]
LAST = ["Kelly", "Santos", "Pham", "Johnson", "Rossi", "Haddad", "Lopez", "Chen", "O'Neil", "Baptiste",
        "Costa", "Murray", "Goldberg", "Mendes", "Brown", "Kowalski", "Doherty", "Reyes", "Silva", "Ward"]
REQUEST_NOTES = ["", "", "", "Rear entrance", "Call before arrival", "Dog on premises", "Use side door",
                 "Lockbox on porch", "Tenant occupied - call ahead", "Contractor on site after 9am"]
MEANINGFUL_NOTES = [
    "Left voicemail for contractor", "Spoke w/ contractor - scheduled", "Contractor not ready - will call back",
    "Owner requested reschedule", "Access issue - left card", "Emailed permit holder", "Scheduled for AM inspection",
]
PLACEHOLDER_NOTES = ["N/A", "TBD", "-", "none"]


# ---------------------------------------------------------------------------
# Request model
# ---------------------------------------------------------------------------
@dataclass
class SynthRequest:
    division: str
    call_dt: datetime | None
    ap_number: str
    address: str
    insp_type: str
    ward: str
    inspector: str
    contact_name: str
    contact_info: str
    request_notes: str
    days: list[date]
    notes: dict[date, str] = field(default_factory=dict)       # note shown from that day onward
    address_on: dict[date, str] = field(default_factory=dict)  # address text override per day
    contact_on: dict[date, str] = field(default_factory=dict)  # contact_info override from that day
    request_notes_on: dict[date, str] = field(default_factory=dict)
    red_override: bool | None = None                           # None = follow the source's 48h rule
    color_mode_on: dict[date, str] = field(default_factory=dict)  # 'theme' -> ambiguous font colour
    blank_inspector_on: set[date] = field(default_factory=set)
    scenario: str = ""

    def note_on(self, d: date) -> str:
        current = ""
        for nd in sorted(self.notes):
            if nd <= d:
                current = self.notes[nd]
        return current

    def latest_override(self, mapping: dict[date, str], d: date, default: str) -> str:
        value = default
        for k in sorted(mapping):
            if k <= d:
                value = mapping[k]
        return value

    def is_red_on(self, d: date) -> bool:
        """The source system's flag: more than 48 clock hours since the call (weekends count)."""
        if self.red_override is not None:
            return self.red_override
        if self.call_dt is None:
            return False
        return datetime.combine(d, time(7, 0)) - self.call_dt > timedelta(hours=48)


# Report days before the storyline (not written out) so earlier calls age out naturally.
_PRE_DAYS: list[date] = [date(2026, 9, 28) + timedelta(days=i) for i in range(5)]  # Mon 9/28 - Fri 10/2
_ALL_DAYS: list[date] = _PRE_DAYS + REPORT_DAYS


def _report_days_after(call_dt: datetime) -> list[date]:
    """Report days whose 07:00 snapshot is after the call (the request is visible from then on)."""
    return [d for d in _ALL_DAYS if datetime.combine(d, time(7, 0)) > call_dt]


class Generator:
    def __init__(self, seed: int = SEED):
        self.rng = random.Random(seed)
        self.used_aps: set[str] = set()
        self.requests: list[SynthRequest] = []
        self.storyline: list[dict] = []

    # -- helpers ----------------------------------------------------------
    def ap(self, division: str) -> str:
        while True:
            n = self.rng.randint(1_000_000, 1_999_999)
            if division == "Building":
                ap = f"ALT{n}"
            elif division == "Electric":
                ap = f"E{n}"
            else:
                # A share of plumbing permits are all-digit with leading zeros (must stay text!)
                ap = f"00{self.rng.randint(100000, 999999)}" if self.rng.random() < 0.25 else f"P{n}"
            if ap not in self.used_aps:
                self.used_aps.add(ap)
                return ap

    def address(self) -> str:
        street, suffix = self.rng.choice(STREETS)
        number = self.rng.randint(2, 1450)
        suffix_txt = suffix if self.rng.random() < 0.3 else ABBR.get(suffix, suffix)
        return f"{number} {street} {suffix_txt}"

    def person(self) -> tuple[str, str]:
        name = f"{self.rng.choice(FIRST)} {self.rng.choice(LAST)}"
        phone = f"(617) 555-01{self.rng.randint(0, 99):02d}"
        return name, phone

    def inspector(self, division: str) -> str:
        names = INSPECTORS[division]
        weights = [INSPECTOR_WEIGHTS.get(n, 1.0) for n in names]
        return self.rng.choices(names, weights=weights, k=1)[0]

    def duration(self) -> int:
        r = self.rng.random()
        if r < 0.25:
            return 1
        if r < 0.48:
            return 2
        if r < 0.66:
            return 3
        if r < 0.87:
            return self.rng.randint(4, 6)
        return self.rng.randint(7, 11)

    # -- background traffic -------------------------------------------------
    def background(self) -> None:
        start, end = date(2026, 9, 21), date(2026, 10, 19)
        d = start
        while d <= end:
            weekend = d.weekday() >= 5
            for division in DIVISIONS:
                n_calls = self.rng.randint(0, 3) if weekend else self.rng.randint(9, 15)
                for _ in range(n_calls):
                    hour = self.rng.uniform(9, 13) if weekend else self.rng.uniform(7.5, 16.5)
                    call_dt = datetime.combine(d, time(0, 0)) + timedelta(
                        hours=int(hour), minutes=int((hour % 1) * 60), seconds=self.rng.randint(0, 59))
                    visible = _report_days_after(call_dt)
                    if not visible:
                        continue
                    days = visible[: self.duration()]
                    # Calls from before the storyline started: only the remaining days show up
                    days = [x for x in days if x >= REPORT_DAYS[0]]
                    if not days:
                        continue
                    name, phone = self.person()
                    req = SynthRequest(
                        division=division, call_dt=call_dt, ap_number=self.ap(division),
                        address=self.address(), insp_type=self.rng.choice(INSP_TYPES[division]),
                        ward=f"{self.rng.randint(1, 22):02d}", inspector=self.inspector(division),
                        contact_name=name, contact_info=phone, request_notes=self.rng.choice(REQUEST_NOTES),
                        days=days,
                    )
                    r = self.rng.random()
                    if r < 0.45:
                        req.notes[self.rng.choice(days)] = self.rng.choice(MEANINGFUL_NOTES)
                        if self.rng.random() < 0.3 and len(days) > 1:
                            later = self.rng.choice(days[1:])
                            req.notes[later] = self.rng.choice(MEANINGFUL_NOTES) + f" ({later:%m/%d})"
                    elif r < 0.55:
                        req.notes[days[0]] = self.rng.choice(PLACEHOLDER_NOTES)
                    self.requests.append(req)
            d += timedelta(days=1)

    # -- scripted scenarios -------------------------------------------------
    def add(self, req: SynthRequest, scenario: str, description: str, expected: dict) -> SynthRequest:
        req.scenario = scenario
        self.used_aps.add(req.ap_number)
        self.requests.append(req)
        self.storyline.append({
            "scenario": scenario,
            "description": description,
            "division": req.division,
            "ap_number": req.ap_number,
            "insp_type": req.insp_type,
            "call_datetime": req.call_dt.isoformat(sep=" ") if req.call_dt else None,
            "report_days": [d.isoformat() for d in req.days],
            "expected": expected,
        })
        return req

    def scripted(self) -> None:
        D = lambda m, d: date(2026, m, d)  # noqa: E731
        T = lambda m, d, h, mi: datetime(2026, m, d, h, mi, 0)  # noqa: E731

        # 1. Spotlight: seen on 4 consecutive reports, notes change, address spelled two ways.
        self.add(SynthRequest(
            division="Building", call_dt=T(10, 5, 9, 14), ap_number="ALT1045821",
            address="88 Dorchester Avenue", insp_type="Rough Frame", ward="07", inspector="Maria Delgado",
            contact_name="Paulo Mendes", contact_info="(617) 555-0142", request_notes="Rear entrance",
            days=[D(10, 6), D(10, 7), D(10, 8), D(10, 9)],
            notes={D(10, 6): "N/A", D(10, 7): "Left voicemail for contractor 10/7",
                   D(10, 8): "Spoke w/ contractor - site not ready, call back Fri",
                   D(10, 9): "Inspection scheduled Tue 10/13 AM"},
            address_on={D(10, 8): "88 DORCHESTER AVE", D(10, 9): "88 Dorchester Ave"},
            contact_on={D(10, 9): "(617) 555-0199"},
            request_notes_on={D(10, 9): ""},   # blank on the last report: must not wipe the old value
        ), "spotlight", "Seen on 4 reports; notes change daily; red from 10/8 but contacted by the due date",
            {"status": "on_time_contacted", "n_reports": 4, "ever_red": True, "mismatch_reason": "excused_by_note",
             "first_note_date": "2026-10-07", "due_date": "2026-10-07"})

        # 2. Same Ap #, different insp_type / call_datetime -> separate IDs.
        common = dict(division="Electric", ap_number="E2051234", address="210 Hyde Park Ave", ward="18",
                      inspector="Linh Tran", contact_name="Keisha Brown", contact_info="(617) 555-0123",
                      request_notes="")
        self.add(SynthRequest(**common, call_dt=T(10, 13, 10, 2), insp_type="Rough Wiring",
                              days=[D(10, 14), D(10, 15)], notes={D(10, 14): "Scheduled for AM inspection"}),
                 "same_ap_type_a", "Same Ap # as two other requests: Rough Wiring called 10/13",
                 {"status": "on_time"})
        self.add(SynthRequest(**common, call_dt=T(10, 13, 10, 2), insp_type="Service Upgrade",
                              days=[D(10, 14)]),
                 "same_ap_type_b", "Same Ap # and call time, different insp_type (Service Upgrade)",
                 {"status": "on_time"})
        self.add(SynthRequest(**common, call_dt=T(10, 15, 8, 40), insp_type="Rough Wiring",
                              days=[D(10, 16), D(10, 19)]),
                 "same_ap_type_c", "Same Ap # and insp_type, called again on 10/15 (different call_datetime); "
                 "still open because Electric's latest report is 10/19 (10/20 email missing)",
                 {"status": "open_not_due"})

        # 3. Red but on time: called Friday before the holiday weekend.
        self.add(SynthRequest(
            division="Building", call_dt=T(10, 9, 10, 30), ap_number="ALT1046677", address="15 Centre St",
            insp_type="Final Building", ward="11", inspector="Kevin O'Brien", contact_name="Mei Chen",
            contact_info="(617) 555-0167", request_notes="Lockbox on porch", days=[D(10, 12), D(10, 13)],
        ), "red_on_time_holiday", "Called Fri 10/9; red on 10/12-10/13 (source counts the holiday weekend); due 10/14",
            {"status": "on_time", "ever_red": True, "mismatch_reason": "calendar_adjustment", "due_date": "2026-10-14"})
        self.add(SynthRequest(
            division="Plumbing/Gas", call_dt=T(10, 10, 11, 5), ap_number="P1077310", address="402 Bennington St",
            insp_type="Gas Pressure Test", ward="01", inspector="Grace Okafor", contact_name="Anna Kowalski",
            contact_info="(617) 555-0108", request_notes="", days=[D(10, 12), D(10, 13), D(10, 14), D(10, 15)],
        ), "weekend_call", "Called Saturday 10/10 -> counts from Tue 10/13 (holiday Monday); due Thu 10/15",
            {"status": "on_time", "due_date": "2026-10-15", "mismatch_reason": "calendar_adjustment"})

        # 4. Late but never red (source missed it).
        self.add(SynthRequest(
            division="Electric", call_dt=T(10, 5, 13, 20), ap_number="E2049876", address="77 Talbot Ave",
            insp_type="Final Electrical", ward="17", inspector="Patrick Walsh", contact_name="Marcus Doherty",
            contact_info="(617) 555-0151", request_notes="Call before arrival",
            days=[D(10, 6), D(10, 7), D(10, 8), D(10, 9), D(10, 12), D(10, 13)], red_override=False,
        ), "late_not_red", "Called 10/5, no note, last seen 10/13, never shown in red",
            {"status": "late", "business_days_late": 3, "mismatch_reason": "missed_by_source"})

        # 5. Late with a note after the due date.
        self.add(SynthRequest(
            division="Plumbing/Gas", call_dt=T(10, 5, 11, 0), ap_number="P1033402", address="9 Savin Hill Ave",
            insp_type="Rough Plumbing", ward="13", inspector="Luis Fernandes", contact_name="Sean Murray",
            contact_info="(617) 555-0177", request_notes="",
            days=[D(10, 6), D(10, 7), D(10, 8), D(10, 9), D(10, 12), D(10, 13), D(10, 14)],
            notes={D(10, 9): "Left voicemail 10/9"},
        ), "late_note_after_due", "Called 10/5, first note 10/9 (after due 10/7), last seen 10/14",
            {"status": "late", "business_days_late": 2})

        # 6. Open-overdue backlog on the final day.
        self.add(SynthRequest(
            division="Building", call_dt=T(10, 13, 9, 0), ap_number="ALT1050012", address="1200 Blue Hill Ave",
            insp_type="Foundation", ward="14", inspector="Thomas Nguyen", contact_name="Yolanda Reyes",
            contact_info="(617) 555-0110", request_notes="Contractor on site after 9am",
            days=[D(10, 14), D(10, 15), D(10, 16), D(10, 19), D(10, 20)],
        ), "open_overdue", "Called 10/13, still outstanding on the final report, no note",
            {"status": "open_overdue", "business_days_late": 3})
        self.add(SynthRequest(
            division="Plumbing/Gas", call_dt=T(10, 7, 14, 45), ap_number="0048812", address="31 Morton St",
            insp_type="Final Plumbing", ward="18", inspector="Erin McCarthy", contact_name="Dmitri Silva",
            contact_info="(617) 555-0133", request_notes="",
            days=[D(10, 8), D(10, 9), D(10, 12), D(10, 13), D(10, 14), D(10, 16), D(10, 19), D(10, 20)],
        ), "open_overdue_leading_zero", "Ap # with leading zeros, open since 10/7 (skips the malformed 10/15 file)",
            {"status": "open_overdue", "ap_number": "0048812"})

        # 7. Missing key field -> review queue.
        self.add(SynthRequest(
            division="Building", call_dt=T(10, 15, 10, 45), ap_number="", address="560 Columbia Rd",
            insp_type="Insulation", ward="15", inspector="Aisha Williams", contact_name="Chloe Santos",
            contact_info="(617) 555-0188", request_notes="", days=[D(10, 16)],
        ), "missing_key", "Ap # blank on the 10/16 Building report -> invalid-row log + review queue",
            {"review_flag": True})

        # 8. Conflict: same Ap #/type/call time as an existing request but a different address.
        self.add(SynthRequest(
            division="Electric", call_dt=T(10, 14, 8, 15), ap_number="E2060001", address="45 Tremont St",
            insp_type="Temporary Service", ward="03", inspector="Sofia Ramirez", contact_name="Ahmed Haddad",
            contact_info="(617) 555-0145", request_notes="", days=[D(10, 15), D(10, 16)],
            notes={D(10, 15): "Spoke w/ contractor - scheduled"},
        ), "conflict_original", "Original request at 45 Tremont St", {"review_flag": False})
        self.add(SynthRequest(
            division="Electric", call_dt=T(10, 14, 8, 15), ap_number="E2060001", address="47 Tremont St",
            insp_type="Temporary Service", ward="03", inspector="Sofia Ramirez", contact_name="Ahmed Haddad",
            contact_info="(617) 555-0145", request_notes="", days=[D(10, 16)],
        ), "conflict_address", "Same Ap #/type/call time but 47 Tremont St -> flagged, not merged",
            {"review_flag": True})

        # 9. Flat Plumbing layout: a blank inspector cell that must NOT be filled down.
        self.add(SynthRequest(
            division="Plumbing/Gas", call_dt=T(10, 8, 9, 30), ap_number="P1088123", address="12 Market St",
            insp_type="Gas Meter Release", ward="22", inspector="Samuel Cohen", contact_name="Rosa Costa",
            contact_info="(617) 555-0119", request_notes="", days=[D(10, 9)], blank_inspector_on={D(10, 9)},
        ), "blank_inspector_flat", "Inspector blank in the flat Plumbing/Gas layout -> left blank, logged",
            {"inspector": ""})

        # 10. Ambiguous (theme) font colour on the Call Date cell.
        self.add(SynthRequest(
            division="Building", call_dt=T(10, 2, 15, 10), ap_number="ALT1039990", address="301 Warren St",
            insp_type="Certificate of Occupancy", ward="12", inspector="Robert Sullivan", contact_name="David Ward",
            contact_info="(617) 555-0104", request_notes="", days=[D(10, 5), D(10, 6), D(10, 7)],
            color_mode_on={D(10, 7): "theme"}, notes={D(10, 6): "Owner requested reschedule"},
        ), "ambiguous_color", "Theme-coloured Call Date cell on 10/7 -> treated as not red and logged",
            {})

    # -- build --------------------------------------------------------------
    def rows_for(self, division: str, day: date) -> list[dict]:
        rows = []
        for req in self.requests:
            if req.division != division or day not in req.days:
                continue
            red = req.is_red_on(day)
            rows.append({
                "inspector": "" if day in req.blank_inspector_on else req.inspector,
                "_group": req.inspector,
                "ward": req.ward,
                "call_dt": req.call_dt,
                "ap_number": req.ap_number,
                "address": req.latest_override(req.address_on, day, req.address),
                "insp_type": req.insp_type,
                "contact_name": req.contact_name,
                "contact_info": req.latest_override(req.contact_on, day, req.contact_info),
                "request_notes": req.latest_override(req.request_notes_on, day, req.request_notes),
                "inspector_notes": req.note_on(day),
                "fulfilled_late": 1 if red else 0,
                "is_red": red,
                "color_mode": req.color_mode_on.get(day, ""),
            })
        rows.sort(key=lambda r: (r["_group"], r["call_dt"] or datetime.min, r["ap_number"]))
        return rows


# ---------------------------------------------------------------------------
# Excel writers (one layout per division, like the real Power BI exports)
# ---------------------------------------------------------------------------
STD_HEADERS = ["Inspector", "Ward", "Call Date", "Ap #", "Address", "Insp Type", "Contact Name",
               "Contact Info", "Request Notes", "Inspector Notes", "FullFilled Late"]
PLUMB_HEADERS = ["Inspector", "Ward", "Call Date/Time", "Application #", "Address", "Inspection Type",
                 "Contact Name", "Contact Info", "Request Notes", "Inspector Notes"]
BOLD = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="FFD9E1F2")
GROUP_FILL = PatternFill("solid", fgColor="FFF2F2F2")
RED_FONT = Font(color="FFFF0000")
THEME_FONT = Font(color=Color(theme=5, tint=0.0))


def _fixed_props(wb: Workbook, day: date) -> None:
    stamp = datetime.combine(day, time(7, 0))
    wb.properties.creator = "Microsoft Power BI (synthetic)"
    wb.properties.created = stamp
    wb.properties.modified = stamp


def _title(ws, division: str, day: date, width: int, extra_rows: list[str]) -> int:
    ws.cell(row=1, column=1, value=SUBJECT[division]).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    r = 2
    for text in extra_rows:
        ws.cell(row=r, column=1, value=text.format(day=day))
        r += 1
    return r


def _call_cell(ws, row: int, col: int, rec: dict, as_text: bool) -> None:
    value = rec["call_dt"]
    if as_text and value is not None:
        value = value.strftime("%m/%d/%Y %I:%M %p")
    cell = ws.cell(row=row, column=col, value=value)
    if not as_text and value is not None:
        cell.number_format = "m/d/yyyy h:mm AM/PM"
    if rec.get("color_mode") == "theme":
        cell.font = THEME_FONT
    elif rec["is_red"]:
        cell.font = RED_FONT


def _std_values(rec: dict) -> list:
    return [rec["inspector"], rec["ward"], None, rec["ap_number"], rec["address"], rec["insp_type"],
            rec["contact_name"], rec["contact_info"], rec["request_notes"] or None,
            rec["inspector_notes"] or None, rec["fulfilled_late"]]


def write_building(path: Path, day: date, rows: list[dict]) -> None:
    """Title rows, header on row 4, inspector cells MERGED vertically per group."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Building"
    r = _title(ws, "Building", day, len(STD_HEADERS), ["Outstanding requests as of {day:%m/%d/%Y} 7:00 AM", ""])
    header_row = r
    for c, h in enumerate(STD_HEADERS, start=1):
        cell = ws.cell(row=header_row, column=c, value=h)
        cell.font, cell.fill = BOLD, HEADER_FILL
    r = header_row + 1
    i = 0
    while i < len(rows):
        group = rows[i]["_group"]
        start = r
        while i < len(rows) and rows[i]["_group"] == group:
            vals = _std_values(rows[i])
            vals[0] = group if r == start else None
            for c, v in enumerate(vals, start=1):
                if c != 3:
                    ws.cell(row=r, column=c, value=v)
            _call_cell(ws, r, 3, rows[i], as_text=False)
            r += 1
            i += 1
        if r - 1 > start:
            ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
            ws.cell(row=start, column=1).alignment = Alignment(vertical="top")
    _fixed_props(wb, day)
    wb.save(path)


def write_electric(path: Path, day: date, rows: list[dict], side_by_side: bool = False) -> None:
    """Header on row 3. Each inspector gets a GROUP HEADER row; data rows beneath have a blank
    inspector cell; groups are separated by a blank row. Optional side-by-side copy."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Electric"
    header_row = _title(ws, "Electric", day, len(STD_HEADERS), [""])
    offsets = [0, len(STD_HEADERS) + 1] if side_by_side else [0]
    for off in offsets:
        for c, h in enumerate(STD_HEADERS, start=1):
            cell = ws.cell(row=header_row, column=c + off, value=h)
            cell.font, cell.fill = BOLD, HEADER_FILL
        r = header_row + 1
        i = 0
        while i < len(rows):
            group = rows[i]["_group"]
            gcell = ws.cell(row=r, column=1 + off, value=group)
            gcell.font, gcell.fill = BOLD, GROUP_FILL
            r += 1
            while i < len(rows) and rows[i]["_group"] == group:
                vals = _std_values(rows[i])
                vals[0] = None
                for c, v in enumerate(vals, start=1):
                    if c != 3:
                        ws.cell(row=r, column=c + off, value=v)
                _call_cell(ws, r, 3 + off, rows[i], as_text=False)
                r += 1
                i += 1
            r += 1  # blank separator row
    _fixed_props(wb, day)
    wb.save(path)


def write_plumbing(path: Path, day: date, rows: list[dict], malformed: bool = False) -> None:
    """Different headers (per-division override), header on row 5, flat layout, text dates."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Plumbing Gas"
    header_row = _title(ws, "Plumbing/Gas", day, len(PLUMB_HEADERS),
                        ["Generated {day:%m/%d/%Y} 7:00 AM", "Filters: Division = Plumbing/Gas", ""])
    headers = list(PLUMB_HEADERS)
    if malformed:
        # The export broke: Call Date and Application # columns came through unnamed.
        headers[2], headers[3] = "Column3", "Column4"
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=c, value=h)
        cell.font, cell.fill = BOLD, HEADER_FILL
    r = header_row + 1
    for rec in rows:
        vals = [rec["inspector"] or None, rec["ward"], None, rec["ap_number"], rec["address"], rec["insp_type"],
                rec["contact_name"], rec["contact_info"], rec["request_notes"] or None, rec["inspector_notes"] or None]
        for c, v in enumerate(vals, start=1):
            if c != 3:
                ws.cell(row=r, column=c, value=v)
        _call_cell(ws, r, 3, rec, as_text=True)
        r += 1
    ws.cell(row=r + 1, column=1, value=f"Total: {len(rows)}").font = BOLD
    _fixed_props(wb, day)
    wb.save(path)


# ---------------------------------------------------------------------------
# Email metadata
# ---------------------------------------------------------------------------
def _email_id(day: date, division: str, n: int = 0) -> str:
    return hashlib.sha1(f"{SEED}|{day}|{division}|{n}".encode()).hexdigest()[:16]


def _email_json(day: date, division: str, filename: str, minute: int = 3, n: int = 0) -> dict:
    received = datetime.combine(day, time(7, minute, 12))
    eid = _email_id(day, division, n)
    return {
        "email_id": eid,
        "thread_id": _email_id(day, division, 99)[:16],
        "from": {"name": SENDER_NAME, "address": SENDER_ADDRESS},
        "to": RECIPIENT,
        "subject": SUBJECT[division],
        "received_at": received.isoformat() + "-04:00",   # EDT in October
        "snippet": (f"Here is your subscription to {SUBJECT[division]}. "
                    f"Outstanding permit inspection requests as of {day:%m/%d/%Y}. "
                    "The attached file contains the report data."),
        "labels": ["INBOX", "ISD-Reports"],
        "attachments": [{
            "attachment_id": "1",
            "filename": f"{SUBJECT[division]}.xlsx",
            "mime_type": XLSX_MIME,
            "file": filename,
        }],
    }


def generate(out_dir: Path = DEFAULT_OUT, seed: int = SEED) -> dict:
    """Generate the whole scripted inbox. Returns the storyline dict (also written to JSON)."""
    out_dir = Path(out_dir)
    if out_dir.exists():
        shutil.rmtree(out_dir)
    out_dir.mkdir(parents=True)
    gen = Generator(seed)
    gen.scripted()      # first, so its Ap # values are reserved before random traffic is drawn
    gen.background()

    counts: dict[str, dict[str, int]] = {}
    for day in REPORT_DAYS:
        day_dir = out_dir / day.isoformat()
        day_dir.mkdir()
        counts[day.isoformat()] = {}
        for division in DIVISIONS:
            rows = gen.rows_for(division, day)
            counts[day.isoformat()][division] = len(rows)
            if division == "Electric" and day == ELECTRIC_MISSING_DAY:
                continue  # the email never arrived
            slug = SLUG[division]
            xlsx = day_dir / f"{slug}.xlsx"
            if division == "Building":
                write_building(xlsx, day, rows)
            elif division == "Electric":
                write_electric(xlsx, day, rows, side_by_side=(day == SIDE_BY_SIDE_ELECTRIC_DAY))
            else:
                write_plumbing(xlsx, day, rows, malformed=(day == MALFORMED_PLUMBING_DAY))
            meta = _email_json(day, division, xlsx.name)
            (day_dir / f"{slug}.json").write_text(json.dumps(meta, indent=2))
            if division == "Building" and day == DUPLICATE_BUILDING_EMAIL_DAY:
                # Power BI re-sent the Building email 28 minutes later; the pipeline keeps the latest.
                resend = _email_json(day, division, xlsx.name, minute=31, n=1)
                (day_dir / f"{slug}_resend.json").write_text(json.dumps(resend, indent=2))

    storyline = {
        "seed": seed,
        "report_days": [d.isoformat() for d in REPORT_DAYS],
        "demo_day": DEMO_DAY.isoformat(),
        "events": {
            "electric_missing": ELECTRIC_MISSING_DAY.isoformat(),
            "malformed_plumbing": MALFORMED_PLUMBING_DAY.isoformat(),
            "duplicate_building_email": DUPLICATE_BUILDING_EMAIL_DAY.isoformat(),
            "side_by_side_electric": SIDE_BY_SIDE_ELECTRIC_DAY.isoformat(),
            "ambiguous_color": AMBIGUOUS_COLOR_DAY.isoformat(),
        },
        "rows_per_day": counts,
        "scenarios": gen.storyline,
    }
    (out_dir / "storyline.json").write_text(json.dumps(storyline, indent=2))
    return storyline


def deliver(scripted_dir: Path, inbox_dir: Path, day: date) -> Path:
    """Copy one scripted day into the mock inbox ("the emails arrive")."""
    src = Path(scripted_dir) / day.isoformat()
    dst = Path(inbox_dir) / day.isoformat()
    if dst.exists():
        shutil.rmtree(dst)
    shutil.copytree(src, dst)
    return dst


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT, help="scripted inbox folder")
    ap.add_argument("--seed", type=int, default=SEED)
    ap.add_argument("--deliver-all", action="store_true", help="also copy every day into the mock inbox")
    args = ap.parse_args(argv)
    story = generate(args.out, args.seed)
    print(f"Wrote {len(story['report_days'])} report days to {args.out}")
    for day, per_div in story["rows_per_day"].items():
        print(f"  {day}: " + ", ".join(f"{d} {n}" for d, n in per_div.items()))
    if args.deliver_all:
        from isd_pipeline.settings import load_settings
        inbox = load_settings().mock_inbox_dir
        for d in REPORT_DAYS:
            deliver(args.out, inbox, d)
        print(f"Delivered all days to {inbox}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
